import os, sqlite3, csv, io, re, json
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file, Response
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import load_workbook, Workbook

BASE=os.path.dirname(os.path.abspath(__file__)); DB=os.path.join(BASE,'infrastructure_bi.db')
app=Flask(__name__); app.secret_key=os.environ.get('BI_SECRET_KEY','local-mvp-change-me')
ROLES=['Admin','Management','Finance','Sales','Project Manager','Business Analyst']
STATUSES=['Planning','Active','On Hold','Completed','Cancelled']
SCHEMAS={
'projects':['project_id','name','customer','location','manager','start_date','expected_date','actual_date','status','completion','contract_value','budget','actual_cost'],
'revenues':['date','project_id','customer','location','product','salesperson','amount'],
'expenses':['date','project_id','category','vendor','amount','payment_status','remarks'],
'customers':['name','segment','location','contact'],'products':['name','category']}

CANONICAL_FIELDS=["project_id","project_name","customer","location","manager","start_date","expected_date","actual_date","status","completion","contract_value","budget","actual_cost","date","product","salesperson","revenue_amount","expense_category","vendor","expense_amount","payment_status","remarks","customer_segment","contact","product_category"]
ALIASES={"project_id":["project id","project code","job id","job no","project no"],"project_name":["project name","project","job name"],"customer":["customer","client","client name","customer name"],"location":["location","city","site","project location"],"manager":["manager","project manager","pm"],"start_date":["start date","project start"],"expected_date":["expected date","completion date","target date","due date"],"actual_date":["actual date","completed date"],"status":["status","project status"],"completion":["completion","completion %","progress","progress %"],"contract_value":["contract value","contract amount","order value","project value"],"budget":["budget","project budget"],"actual_cost":["actual cost","project cost","actual project cost"],"date":["date","transaction date","invoice date","revenue date","expense date"],"product":["product","product/service","service","item","product name"],"salesperson":["salesperson","sales person","sales executive","executive"],"revenue_amount":["revenue","revenue amount","sales","sales amount","invoice amount"],"expense_category":["expense category","category","cost category"],"vendor":["vendor","supplier"],"expense_amount":["expense amount","cost amount","expense","cost"],"payment_status":["payment status","paid status"],"remarks":["remarks","remark","notes","comments"],"customer_segment":["segment","customer segment","customer category","category"],"contact":["contact","phone","mobile","contact number"],"product_category":["product category","service category"]}
def nh(x):
 x=re.sub(r'([a-z0-9])([A-Z])',r'\1 \2',str(x or ''))
 return re.sub(r'[^a-z0-9]+',' ',x.lower()).strip()
def auto_map(headers):
 out={}
 for h in headers:
  n=nh(h)
  for f,als in ALIASES.items():
   if n==nh(f) or n in [nh(a) for a in als]: out[h]=f; break
 return out
def validate_rows(rows,mapping):
 errs=[]; warns=[]
 if not rows: errs.append('CSV contains no data rows.')
 if not any(x in mapping.values() for x in ['project_id','project_name']): warns.append('No project ID/name mapped; Projects may remain empty.')
 if 'customer' not in mapping.values(): warns.append('Customer column not mapped; Customer analytics will be limited.')
 if not any(x in mapping.values() for x in ['revenue_amount','expense_amount','actual_cost','budget','contract_value']): warns.append('No financial amount column mapped; financial KPIs may be zero.')
 for i,r in enumerate(rows[:1000],2):
  for f in ['completion','contract_value','budget','actual_cost','revenue_amount','expense_amount']:
   src=next((h for h,v in mapping.items() if v==f),None)
   if src and str(r.get(src,'')).strip():
    try: float(str(r[src]).replace(',','').replace('%',''))
    except: errs.append(f'Row {i}: {src} is not numeric.')
    if len(errs)>=20:return errs,warns
 return errs,warns
def load_unified(rows,mapping):
 c=db(); counts={x:0 for x in ['projects','revenues','expenses','customers','products']}; cust=set(); prod=set()
 def get(r,f):
  h=next((h for h,v in mapping.items() if v==f),None); return r.get(h,'') if h else ''
 for idx,r in enumerate(rows,2):
  pid=str(get(r,'project_id')).strip() or ('IMP-'+str(idx)); pn=str(get(r,'project_name')).strip(); customer=str(get(r,'customer')).strip(); product=str(get(r,'product')).strip()
  if pn or 'project_id' in mapping:
   c.execute('INSERT OR IGNORE INTO projects(project_id,name,customer,location,manager,start_date,expected_date,actual_date,status,completion,contract_value,budget,actual_cost) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(pid,pn or pid,customer,get(r,'location'),get(r,'manager'),get(r,'start_date'),get(r,'expected_date'),get(r,'actual_date'),get(r,'status') or 'Active',num(get(r,'completion')),num(get(r,'contract_value')),num(get(r,'budget')),num(get(r,'actual_cost')))); counts['projects']+=1
  if customer and customer.lower() not in cust:
   c.execute('INSERT OR IGNORE INTO customers(name,segment,location,contact) VALUES(?,?,?,?)',(customer,get(r,'customer_segment'),get(r,'location'),get(r,'contact'))); cust.add(customer.lower()); counts['customers']+=1
  if product and product.lower() not in prod:
   c.execute('INSERT OR IGNORE INTO products(name,category) VALUES(?,?)',(product,get(r,'product_category'))); prod.add(product.lower()); counts['products']+=1
  rv=get(r,'revenue_amount')
  if str(rv).strip(): c.execute('INSERT INTO revenues(date,project_id,customer,location,product,salesperson,amount) VALUES(?,?,?,?,?,?,?)',(get(r,'date'),pid,customer,get(r,'location'),product,get(r,'salesperson'),num(rv))); counts['revenues']+=1
  ex=get(r,'expense_amount')
  if str(ex).strip(): c.execute('INSERT INTO expenses(date,project_id,category,vendor,amount,payment_status,remarks) VALUES(?,?,?,?,?,?,?)',(get(r,'date'),pid,get(r,'expense_category'),get(r,'vendor'),num(ex),get(r,'payment_status'),get(r,'remarks'))); counts['expenses']+=1
 c.commit(); c.close(); return counts
def db():
 c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; return c
def init():
 c=db(); q=c.execute
 q("""CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY AUTOINCREMENT,username TEXT UNIQUE,password_hash TEXT,role TEXT,active INTEGER DEFAULT 1)""")
 q("""CREATE TABLE IF NOT EXISTS projects(id INTEGER PRIMARY KEY AUTOINCREMENT,project_id TEXT UNIQUE,name TEXT,customer TEXT,location TEXT,manager TEXT,start_date TEXT,expected_date TEXT,actual_date TEXT,status TEXT,completion REAL DEFAULT 0,contract_value REAL DEFAULT 0,budget REAL DEFAULT 0,actual_cost REAL DEFAULT 0)""")
 q("""CREATE TABLE IF NOT EXISTS revenues(id INTEGER PRIMARY KEY AUTOINCREMENT,date TEXT,project_id TEXT,customer TEXT,location TEXT,product TEXT,salesperson TEXT,amount REAL DEFAULT 0)""")
 q("""CREATE TABLE IF NOT EXISTS expenses(id INTEGER PRIMARY KEY AUTOINCREMENT,date TEXT,project_id TEXT,category TEXT,vendor TEXT,amount REAL DEFAULT 0,payment_status TEXT,remarks TEXT)""")
 q("""CREATE TABLE IF NOT EXISTS customers(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT UNIQUE,segment TEXT,location TEXT,contact TEXT)""")
 q("""CREATE TABLE IF NOT EXISTS products(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT UNIQUE,category TEXT)""")
 if not q('SELECT 1 FROM users LIMIT 1').fetchone(): q('INSERT INTO users(username,password_hash,role) VALUES(?,?,?)',('admin',generate_password_hash('admin123'),'Admin'))
 c.commit(); c.close()
def auth(f):
 @wraps(f)
 def w(*a,**k): return f(*a,**k) if 'uid' in session else redirect(url_for('login'))
 return w
def admin(f):
 @wraps(f)
 def w(*a,**k):
  if 'uid' not in session:return redirect(url_for('login'))
  if session.get('role')!='Admin': flash('Admin access required.','danger'); return redirect(url_for('dashboard'))
  return f(*a,**k)
 return w
def num(v):
 try:return float(v or 0)
 except:return 0
@app.context_processor
def ctx(): return {'current_user':session.get('username'),'current_role':session.get('role')}
@app.route('/')
def home():
 if 'uid' not in session:return redirect(url_for('login'))
 c=db(); n=sum(c.execute(f'SELECT COUNT(*)x FROM {t}').fetchone()['x'] for t in SCHEMAS); c.close()
 return redirect(url_for('startup_import') if n==0 else url_for('dashboard'))
@app.route('/login',methods=['GET','POST'])
def login():
 if request.method=='POST':
  c=db(); u=c.execute('SELECT * FROM users WHERE username=? AND active=1',(request.form['username'],)).fetchone(); c.close()
  if u and check_password_hash(u['password_hash'],request.form['password']): session.update(uid=u['id'],username=u['username'],role=u['role']); return redirect(url_for('dashboard'))
  flash('Invalid username or password.','danger')
 return render_template('login.html')
@app.route('/logout')
def logout():session.clear();return redirect(url_for('login'))
@app.route('/startup-import',methods=['GET','POST'])
@auth
def startup_import():
 if request.method=='POST':
  f=request.files.get('file')
  if not f or not f.filename.lower().endswith('.csv'): flash('Please select a CSV file.','danger'); return redirect(url_for('startup_import'))
  try:
   reader=csv.DictReader(io.TextIOWrapper(f.stream,encoding='utf-8-sig',newline='')); headers=reader.fieldnames or []; rows=list(reader); mapping=auto_map(headers); errs,warns=validate_rows(rows,mapping)
   if errs: return render_template('startup_import.html',headers=headers,mapping=mapping,rows=rows[:8],all_rows=rows,errors=errs,warnings=warns,step='map',fields=CANONICAL_FIELDS)
   return render_template('startup_import.html',headers=headers,mapping=mapping,rows=rows[:8],errors=[],warnings=warns,step='map',all_rows=rows,fields=CANONICAL_FIELDS)
  except Exception as e: flash('CSV read error: '+str(e),'danger')
 return render_template('startup_import.html',headers=[],mapping={},rows=[],errors=[],warnings=[],step='upload',fields=CANONICAL_FIELDS)
@app.route('/startup-import/load',methods=['POST'])
@auth
def startup_import_load():
 try:
  headers=request.form.getlist('headers'); rows=json.loads(request.form['rows_json']); mapping={h:request.form.get('map_'+str(i),'') for i,h in enumerate(headers)}; mapping={h:v for h,v in mapping.items() if v}; errs,warns=validate_rows(rows,mapping)
  if errs: flash('Validation failed: '+' | '.join(errs[:5]),'danger'); return redirect(url_for('startup_import'))
  load_unified(rows,mapping); flash('CSV validated, mapped, loaded into SQLite, and analytics populated.','success'); return redirect(url_for('dashboard'))
 except Exception as e: flash('Import failed: '+str(e),'danger'); return redirect(url_for('startup_import'))

@app.route('/dashboard')
@auth
def dashboard():
 c=db(); revenue=c.execute('SELECT COALESCE(SUM(amount),0)x FROM revenues').fetchone()['x']; expenses=c.execute('SELECT COALESCE(SUM(amount),0)x FROM expenses').fetchone()['x']; cost=c.execute('SELECT COALESCE(SUM(actual_cost),0)x FROM projects').fetchone()['x']; budget=c.execute('SELECT COALESCE(SUM(budget),0)x FROM projects').fetchone()['x']; contract=c.execute('SELECT COALESCE(SUM(contract_value),0)x FROM projects').fetchone()['x']; active=c.execute("SELECT COUNT(*)x FROM projects WHERE status='Active'").fetchone()['x']; completed=c.execute("SELECT COUNT(*)x FROM projects WHERE status='Completed'").fetchone()['x']; monthly=c.execute("SELECT substr(date,1,7)m,SUM(amount)total FROM revenues GROUP BY m ORDER BY m DESC LIMIT 12").fetchall(); status=c.execute('SELECT status,COUNT(*)total FROM projects GROUP BY status').fetchall(); recent=c.execute('SELECT * FROM projects ORDER BY id DESC LIMIT 8').fetchall(); loc=c.execute('SELECT location,SUM(amount)total FROM revenues GROUP BY location ORDER BY total DESC LIMIT 8').fetchall(); c.close(); profit=revenue-expenses-cost
 return render_template('dashboard.html',revenue=revenue,expenses=expenses,cost=cost,budget=budget,contract=contract,active=active,completed=completed,profit=profit,margin=(profit/revenue*100 if revenue else 0),monthly=monthly,status=status,recent=recent,locations=loc)
@app.route('/projects',methods=['GET','POST'])
@auth
def projects():
 c=db()
 if request.method=='POST':
  try:c.execute('INSERT INTO projects(project_id,name,customer,location,manager,start_date,expected_date,actual_date,status,completion,contract_value,budget,actual_cost) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(request.form['project_id'],request.form['name'],request.form.get('customer'),request.form.get('location'),request.form.get('manager'),request.form.get('start_date'),request.form.get('expected_date'),request.form.get('actual_date'),request.form.get('status','Planning'),num(request.form.get('completion')),num(request.form.get('contract_value')),num(request.form.get('budget')),num(request.form.get('actual_cost'))));c.commit();flash('Project added.','success')
  except sqlite3.IntegrityError:flash('Project ID already exists.','danger')
 q=request.args.get('q',''); st=request.args.get('status',''); sql='SELECT * FROM projects WHERE 1=1'; p=[]
 if q:sql+=' AND (project_id LIKE ? OR name LIKE ? OR customer LIKE ? OR location LIKE ?)';p += [f'%{q}%']*4
 if st:sql+=' AND status=?';p.append(st)
 rows=c.execute(sql+' ORDER BY id DESC',p).fetchall();c.close();return render_template('projects.html',rows=rows,q=q,status=st,statuses=STATUSES)
@app.route('/revenue',methods=['GET','POST'])
@auth
def revenue():
 c=db()
 if request.method=='POST':c.execute('INSERT INTO revenues(date,project_id,customer,location,product,salesperson,amount) VALUES(?,?,?,?,?,?,?)',(request.form.get('date'),request.form.get('project_id'),request.form.get('customer'),request.form.get('location'),request.form.get('product'),request.form.get('salesperson'),num(request.form.get('amount'))));c.commit();flash('Revenue added.','success')
 q=request.args.get('q','');rows=c.execute('SELECT * FROM revenues WHERE project_id LIKE ? OR customer LIKE ? OR location LIKE ? OR product LIKE ? OR salesperson LIKE ? ORDER BY date DESC,id DESC',(f'%{q}%',)*5).fetchall();total=sum(r['amount'] for r in rows);c.close();return render_template('revenue.html',rows=rows,total=total,q=q)
@app.route('/expenses',methods=['GET','POST'])
@auth
def expenses():
 c=db()
 if request.method=='POST':c.execute('INSERT INTO expenses(date,project_id,category,vendor,amount,payment_status,remarks) VALUES(?,?,?,?,?,?,?)',(request.form.get('date'),request.form.get('project_id'),request.form.get('category'),request.form.get('vendor'),num(request.form.get('amount')),request.form.get('payment_status'),request.form.get('remarks')));c.commit();flash('Expense added.','success')
 q=request.args.get('q','');rows=c.execute('SELECT * FROM expenses WHERE project_id LIKE ? OR category LIKE ? OR vendor LIKE ? OR payment_status LIKE ? ORDER BY date DESC,id DESC',(f'%{q}%',)*4).fetchall();total=sum(r['amount'] for r in rows);c.close();return render_template('expenses.html',rows=rows,total=total,q=q)
@app.route('/customers',methods=['GET','POST'])
@auth
def customers():
 c=db()
 if request.method=='POST':
  try:c.execute('INSERT INTO customers(name,segment,location,contact) VALUES(?,?,?,?)',(request.form['name'],request.form.get('segment'),request.form.get('location'),request.form.get('contact')));c.commit();flash('Customer added.','success')
  except sqlite3.IntegrityError:flash('Customer already exists.','danger')
 q=request.args.get('q','');rows=c.execute('SELECT * FROM customers WHERE name LIKE ? OR segment LIKE ? OR location LIKE ? ORDER BY name',(f'%{q}%',)*3).fetchall();analytics=c.execute('SELECT c.name,c.segment,c.location,COALESCE(SUM(r.amount),0)revenue,COUNT(DISTINCT r.project_id)projects FROM customers c LEFT JOIN revenues r ON lower(r.customer)=lower(c.name) GROUP BY c.id ORDER BY revenue DESC').fetchall();c.close();return render_template('customers.html',rows=rows,analytics=analytics,q=q)
@app.route('/products',methods=['GET','POST'])
@auth
def products():
 c=db()
 if request.method=='POST':
  try:c.execute('INSERT INTO products(name,category) VALUES(?,?)',(request.form['name'],request.form.get('category')));c.commit();flash('Product/service added.','success')
  except sqlite3.IntegrityError:flash('Product/service already exists.','danger')
 rows=c.execute('SELECT * FROM products ORDER BY name').fetchall();analytics=c.execute('SELECT p.name,p.category,COALESCE(SUM(r.amount),0)revenue,COUNT(r.id)transactions FROM products p LEFT JOIN revenues r ON lower(r.product)=lower(p.name) GROUP BY p.id ORDER BY revenue DESC').fetchall();c.close();return render_template('products.html',rows=rows,analytics=analytics)
@app.route('/reports')
@auth
def reports():
 c=db();by_project=c.execute("SELECT p.project_id,p.name,p.budget,p.actual_cost,COALESCE((SELECT SUM(amount) FROM revenues r WHERE r.project_id=p.project_id),0)revenue,p.budget-p.actual_cost variance FROM projects p ORDER BY revenue DESC").fetchall();loc=c.execute('SELECT location,SUM(amount)revenue FROM revenues GROUP BY location ORDER BY revenue DESC').fetchall();cust=c.execute('SELECT customer,SUM(amount)revenue FROM revenues GROUP BY customer ORDER BY revenue DESC').fetchall();prod=c.execute('SELECT product,SUM(amount)revenue FROM revenues GROUP BY product ORDER BY revenue DESC').fetchall();mon=c.execute("SELECT substr(date,1,7)month,SUM(amount)revenue FROM revenues GROUP BY month ORDER BY month").fetchall();c.close();return render_template('reports.html',by_project=by_project,loc=loc,cust=cust,prod=prod,mon=mon)
@app.route('/import',methods=['GET','POST'])
@auth
def imp():
 if request.method=='POST':
  target=request.form['target'];f=request.files.get('file')
  try:
   if f.filename.lower().endswith('.csv'): data=list(csv.DictReader(io.TextIOWrapper(f.stream,encoding='utf-8-sig',newline='')))
   else:
    wb=load_workbook(f.stream,data_only=True,read_only=True);v=list(wb.active.values);heads=[str(x).strip() if x is not None else '' for x in v[0]];data=[dict(zip(heads,row)) for row in v[1:]]
   cols={'projects':['project_id','name','customer','location','manager','start_date','expected_date','actual_date','status','completion','contract_value','budget','actual_cost'],'revenues':['date','project_id','customer','location','product','salesperson','amount'],'expenses':['date','project_id','category','vendor','amount','payment_status','remarks'],'customers':['name','segment','location','contact'],'products':['name','category']}[target];c=db();n=0
   for row in data:
    vals=[num(row.get(x)) if x in ('completion','contract_value','budget','actual_cost','amount') else (row.get(x) or '') for x in cols];c.execute(f"INSERT OR IGNORE INTO {target}({','.join(cols)}) VALUES({','.join(['?']*len(cols))})",vals);n+=1
   c.commit();c.close();flash(f'Imported {n} rows.','success')
  except Exception as e:flash('Import failed: '+str(e),'danger')
 return render_template('import.html')
@app.route('/export/<target>')
@auth
def export(target):
 cols={'projects':['project_id','name','customer','location','manager','start_date','expected_date','actual_date','status','completion','contract_value','budget','actual_cost'],'revenues':['date','project_id','customer','location','product','salesperson','amount'],'expenses':['date','project_id','category','vendor','amount','payment_status','remarks'],'customers':['name','segment','location','contact'],'products':['name','category']}[target];c=db();rows=c.execute('SELECT * FROM '+target).fetchall();c.close();fmt=request.args.get('format','xlsx')
 if fmt=='csv':
  s=io.StringIO();w=csv.writer(s);w.writerow(cols);[w.writerow([r[x] for x in cols]) for r in rows];return Response('\ufeff'+s.getvalue(),mimetype='text/csv',headers={'Content-Disposition':f'attachment; filename={target}.csv'})
 wb=Workbook();ws=wb.active;ws.append(cols);[ws.append([r[x] for x in cols]) for r in rows];b=io.BytesIO();wb.save(b);b.seek(0);return send_file(b,as_attachment=True,download_name=target+'.xlsx')
@app.route('/users',methods=['GET','POST'])
@admin
def users():
 c=db()
 if request.method=='POST':
  try:c.execute('INSERT INTO users(username,password_hash,role) VALUES(?,?,?)',(request.form['username'],generate_password_hash(request.form['password']),request.form['role']));c.commit();flash('User created.','success')
  except sqlite3.IntegrityError:flash('Username already exists.','danger')
 rows=c.execute('SELECT id,username,role,active FROM users ORDER BY username').fetchall();c.close();return render_template('users.html',rows=rows,roles=ROLES)
if __name__=='__main__':init();app.run(host='127.0.0.1',port=5000,debug=False)
